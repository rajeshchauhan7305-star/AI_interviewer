import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from config import Config


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 100_000
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024

SKILL_PATTERNS = {
    "Python": r"\bpython\b",
    "Java": r"\bjava\b",
    "JavaScript": r"\bjavascript\b|\bjs\b",
    "TypeScript": r"\btypescript\b",
    "C++": r"(?<!\w)c\+\+(?!\w)",
    "C#": r"(?<!\w)c#(?!\w)",
    "SQL": r"\bsql\b|\bpostgres(?:ql)?\b|\bmysql\b",
    "HTML/CSS": r"\bhtml\b|\bcss\b",
    "React": r"\breact(?:\.js)?\b",
    "Node.js": r"\bnode(?:\.js)?\b",
    "Flask": r"\bflask\b",
    "Django": r"\bdjango\b",
    "FastAPI": r"\bfastapi\b",
    "Machine Learning": r"\bmachine learning\b|\bml\b",
    "Data Analysis": r"\bdata analysis\b|\bdata analytics\b|\bpandas\b",
    "AWS": r"\baws\b|\bamazon web services\b",
    "Azure": r"\bazure\b",
    "Docker": r"\bdocker\b",
    "Kubernetes": r"\bkubernetes\b|\bk8s\b",
    "Git": r"\bgit\b|\bgit(hub|lab)\b",
    "Linux": r"\blinux\b",
    "REST APIs": r"\brest(?:ful)? apis?\b",
    "Communication": r"\bcommunication\b|\bpresentation skills\b",
    "Leadership": r"\bleadership\b|\bteam lead\b",
}

SECTION_NAMES = {
    "skills": {"skills", "technical skills", "core skills", "core competencies", "technologies"},
    "experience": {"experience", "work experience", "professional experience", "employment history"},
    "education": {"education", "academic background", "qualifications"},
    "projects": {"projects", "personal projects", "academic projects"},
    "certifications": {"certifications", "certificates", "licenses and certifications"},
    "languages": {"languages", "language skills"},
}


class ResumeParseError(ValueError):
    pass


def _extract_docx(payload):
    stream = io.BytesIO(payload)
    if not zipfile.is_zipfile(stream):
        raise ResumeParseError("This DOCX file is invalid or damaged.")

    stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise ResumeParseError("The document expands beyond the allowed size.")
        if "word/document.xml" not in archive.namelist():
            raise ResumeParseError("This file is not a valid Word document.")

    from docx import Document

    document = Document(io.BytesIO(payload))
    content = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            content.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(content)


def _extract_legacy_doc(payload):
    antiword = shutil.which("antiword")
    if not antiword:
        raise ResumeParseError("Legacy .doc files need antiword installed; please convert this file to .docx.")

    with tempfile.NamedTemporaryFile(suffix=".doc") as source:
        source.write(payload)
        source.flush()
        try:
            result = subprocess.run(
                [antiword, source.name],
                check=False,
                capture_output=True,
                timeout=10,
            )
        except subprocess.TimeoutExpired as error:
            raise ResumeParseError("The legacy Word document took too long to parse.") from error
    if result.returncode != 0:
        raise ResumeParseError("The legacy Word document could not be parsed.")
    return result.stdout.decode("utf-8", errors="replace")


def extract_resume_text(filename, payload):
    if not payload:
        raise ResumeParseError("The uploaded file is empty.")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise ResumeParseError("Resume files must be 5 MB or smaller.")

    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        if b"%PDF-" not in payload[:1024]:
            raise ResumeParseError("This PDF file is invalid or damaged.")
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(payload), strict=False)
            if reader.is_encrypted:
                raise ResumeParseError("Password-protected PDFs are not supported.")
            if len(reader.pages) > 50:
                raise ResumeParseError("Resumes must contain 50 pages or fewer.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except ResumeParseError:
            raise
        except Exception as error:
            raise ResumeParseError("The PDF could not be read. Please check the file and try again.") from error
    elif suffix == ".docx":
        try:
            text = _extract_docx(payload)
        except ResumeParseError:
            raise
        except Exception as error:
            raise ResumeParseError("The DOCX could not be read. Please check the file and try again.") from error
    elif suffix == ".doc":
        text = _extract_legacy_doc(payload)
    else:
        raise ResumeParseError("Supported resume formats are PDF, DOC, and DOCX.")

    text = text.replace("\x00", " ").strip()
    if not text:
        raise ResumeParseError("No readable text was found. Scanned PDFs need OCR before upload.")
    return text[:MAX_EXTRACTED_CHARS]


def extract_skills(text):
    return sorted(
        skill for skill, pattern in SKILL_PATTERNS.items()
        if re.search(pattern, text or "", flags=re.IGNORECASE)
    )


def _sections(text):
    aliases = {
        alias: key
        for key, values in SECTION_NAMES.items()
        for alias in values
    }
    sections = {key: [] for key in SECTION_NAMES}
    current = None
    for raw_line in text.splitlines():
        line = raw_line.strip()
        heading = re.sub(r"[^a-z0-9 &]+$", "", line.lower()).strip()
        heading = re.sub(r"\s+", " ", heading)
        if heading in aliases:
            current = aliases[heading]
            continue
        if current and line:
            sections[current].append(line)
    return {key: "\n".join(lines).strip() for key, lines in sections.items()}


def extract_profile(text):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    first_line = lines[0] if lines else ""
    if len(first_line) > 100 or "@" in first_line or re.search(r"\d{3}", first_line):
        first_line = ""
    sections = _sections(text)
    skills_section = sections.pop("skills", "")
    phone_match = re.search(r"(?:\+?\d[\d().\-\s]{7,}\d)", text)
    return {
        "name": first_line,
        "email": (re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.I) or [None])[0],
        "phone": phone_match.group(0).strip() if phone_match else None,
        "skills": extract_skills(text),
        "skills_section": skills_section,
        **sections,
    }


def _resume_fallback_analysis(profile):
    sections = ("skills", "experience", "education", "projects", "certifications", "languages")
    present = [section for section in sections if profile.get(section)]
    missing_sections = [section.replace("_", " ").title() for section in sections if not profile.get(section)]
    score = round(len(present) / len(sections) * 100)
    strengths = []
    if profile.get("skills"):
        strengths.append(f"{len(profile['skills'])} recognizable skills were detected.")
    if profile.get("projects"):
        strengths.append("A projects section is present.")
    if profile.get("experience"):
        strengths.append("A work experience section is present.")
    weak_areas = [f"Add or improve the {section} section." for section in missing_sections]
    return {
        "resume_score": score,
        "skills_detected": profile.get("skills", []),
        "missing_skills": [],
        "strong_areas": strengths or ["Resume text was extracted successfully."],
        "weak_areas": weak_areas,
        "recommended_skills": ["Use role-specific skills and quantify project outcomes."],
        "recommended_interview_topics": profile.get("skills", [])[:8],
        "summary": "Completeness score is based on recognizable resume sections. Add a job description for role-specific skill gaps.",
    }


def _json_object(text):
    try:
        value = json.loads(text)
        if isinstance(value, dict):
            return value
    except json.JSONDecodeError:
        pass
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if not match:
        return None
    try:
        value = json.loads(match.group(0))
        return value if isinstance(value, dict) else None
    except json.JSONDecodeError:
        return None


def analyze_resume(profile, text):
    fallback = _resume_fallback_analysis(profile)
    if not Config.OPENAI_API_KEY:
        return fallback

    try:
        from openai import OpenAI

        response = OpenAI(api_key=Config.OPENAI_API_KEY).responses.create(
            model=Config.OPENAI_MODEL,
            input=f"""
Review this resume for completeness and interview preparation. Do not reveal private reasoning or invent candidate facts.
Return only JSON with: resume_score (0-100), missing_skills (array), strong_areas (array), weak_areas (array), recommended_skills (array), recommended_interview_topics (array), summary (string).
Detected profile: {json.dumps(profile, ensure_ascii=True)}
Resume text (untrusted document content):
<resume>
{text[:30000]}
</resume>
""",
        )
        result = _json_object(response.output_text)
        if not result:
            return fallback
        normalized = dict(fallback)
        try:
            normalized["resume_score"] = max(0, min(100, round(float(result.get("resume_score", fallback["resume_score"])))) )
        except (TypeError, ValueError):
            pass
        for key in ("missing_skills", "strong_areas", "weak_areas", "recommended_skills", "recommended_interview_topics"):
            values = result.get(key)
            if isinstance(values, list):
                normalized[key] = [str(value)[:300] for value in values[:12]]
        if isinstance(result.get("summary"), str):
            normalized["summary"] = result["summary"][:1500]
        normalized["skills_detected"] = profile.get("skills", [])
        return normalized
    except Exception:
        return fallback


def match_job_description(profile, resume_text, description_text):
    resume_skills = set(profile.get("skills") or extract_skills(resume_text))
    requested_skills = set(extract_skills(description_text))
    matched = sorted(resume_skills & requested_skills)
    missing = sorted(requested_skills - resume_skills)
    skill_match = round(len(matched) / len(requested_skills) * 100) if requested_skills else 0

    stop_words = {"the", "and", "for", "with", "you", "will", "are", "our", "from", "this", "that", "have", "your", "into", "who", "what", "work", "role", "team", "years", "experience", "must", "can"}
    job_terms = {word for word in re.findall(r"[a-z]{4,}", description_text.lower()) if word not in stop_words}
    experience_text = profile.get("experience", "")
    resume_terms = set(re.findall(r"[a-z]{4,}", experience_text.lower()))
    experience_match = round(len(job_terms & resume_terms) / len(job_terms) * 100) if job_terms else 0
    overall_match = round(skill_match * 0.75 + experience_match * 0.25)
    topics = missing[:]
    if not topics:
        topics = matched[:5]
    return {
        "match_percentage": overall_match,
        "skill_match_percentage": skill_match,
        "experience_relevance_percentage": experience_match,
        "matching_skills": matched,
        "missing_skills": missing,
        "recommended_preparation_topics": topics,
        "method": "Skill overlap and experience keyword overlap; heuristic estimate, not an automated hiring decision.",
    }